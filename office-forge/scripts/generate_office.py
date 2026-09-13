#!/usr/bin/env python3
"""
Office Forge Generator
Scaffolds Python generation scripts for building clean .xlsx and .pptx files on Windows.
"""

import sys
from pathlib import Path

# Force UTF-8 on Windows
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

XLSX_BUILDER = '''#!/usr/bin/env python3
"""Builds formatted Excel workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def build_excel(output_path: str = "output.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "Category"
    ws["B1"] = "Value"
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["B1"].font = Font(bold=True, color="FFFFFF")
    fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws["A1"].fill = fill
    ws["B1"].fill = fill

    data = [("Item Alpha", 100), ("Item Beta", 250), ("Item Gamma", 420)]
    for row_idx, (cat, val) in enumerate(data, start=2):
        ws[f"A{row_idx}"] = cat
        ws[f"B{row_idx}"] = val
        ws[f"B{row_idx}"].font = Font(color="0000FF")

    ws[f"A{len(data)+2}"] = "Total"
    ws[f"B{len(data)+2}"] = f"=SUM(B2:B{len(data)+1})"
    ws[f"A{len(data)+2}"].font = Font(bold=True)
    ws[f"B{len(data)+2}"].font = Font(bold=True)

    wb.save(output_path)
    print(f"[OK] Workbook saved to: {output_path}")

if __name__ == "__main__":
    build_excel()
'''

PPTX_BUILDER = '''#!/usr/bin/env python3
"""Builds 16:9 widescreen PowerPoint presentation."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def build_presentation(output_path: str = "deck.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.0))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = "Executive Briefing"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    body_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.5))
    btf = body_box.text_frame
    bp = btf.paragraphs[0]
    bp.text = "• Strategic Objective 1: High velocity execution"
    bp.font.size = Pt(16)

    prs.save(output_path)
    print(f"[OK] Presentation saved to: {output_path}")

if __name__ == "__main__":
    build_presentation()
'''

def scaffold_builder(doc_type: str = "xlsx", out_file: str = "build_doc.py"):
    target = Path(out_file).resolve()
    code = XLSX_BUILDER if doc_type.lower() == "xlsx" else PPTX_BUILDER
    target.write_text(code, encoding="utf-8")
    print(f"[OK] Generated {doc_type.upper()} builder at: {target}")

if __name__ == "__main__":
    dtype = sys.argv[1] if len(sys.argv) > 1 else "xlsx"
    out = sys.argv[2] if len(sys.argv) > 2 else f"build_{dtype}.py"
    scaffold_builder(dtype, out)
